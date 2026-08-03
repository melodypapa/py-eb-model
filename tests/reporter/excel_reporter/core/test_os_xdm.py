"""
Os Excel Reporter Tests.

Implements:
    - TC_UNIT_REPORTER_00016: Os Excel Reporter - File Creation
    - TC_UNIT_REPORTER_00017: Os Excel Reporter - Edge Cases

Implements: SWR_REPORTER_00002
"""
import os
import tempfile
from openpyxl import load_workbook
from eb_model.models.core.abstract import EcucRefType
from eb_model.models.core.eb_doc import EBModel
from eb_model.models.core.os_xdm import (
    Os, OsSpinlock, OsTask, OsApplication, OsOS, OsHooks,
    OsMicrokernel, MkMemoryProtection, MkMemoryRegion,
    OsScheduleTable, OsIsr, OsCounter, OsPeripheralArea
)
from eb_model.reporter.excel_reporter.core.os_xdm import OsXdmXlsWriter


class TestOsXdmXlsWriter:

    def test_write_creates_excel_file(self):
        """Test that write() creates a valid Excel file.

        Implements: TC_UNIT_REPORTER_00016
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            writer.write(filename, doc, options={"skip_os_task": True})

            assert os.path.exists(filename)
            assert os.path.getsize(filename) > 0

            wb = load_workbook(filename)
            # OS reporter creates multiple sheets
            assert len(wb.sheetnames) > 0
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_creates_expected_sheets(self):
        """Test that write() creates expected sheets.

        Implements: TC_UNIT_REPORTER_00016
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            writer.write(filename, doc, options={"skip_os_task": True})

            wb = load_workbook(filename)

            # Check for expected sheets
            expected_sheets = ["OsApplications", "OsIsr"]
            for sheet_name in expected_sheets:
                assert sheet_name in wb.sheetnames, f"Sheet {sheet_name} not found"

            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_spinlock_multiple_accessing_apps(self):
        """
        Test spinlock sheet with multiple accessing applications (wrap text).

        Implements: TC_UNIT_REPORTER_00017
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            spinlock = OsSpinlock(os_mod, "TestSpinlock")
            spinlock.setOsSpinlockLockMethod("STANDARD")
            spinlock.addOsSpinlockAccessingApplication(EcucRefType("ASPath:/Os/App1"))
            spinlock.addOsSpinlockAccessingApplication(EcucRefType("ASPath:/Os/App2"))
            os_mod.addOsSpinlock(spinlock)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsSpinlock" in wb.sheetnames
            sheet = wb["OsSpinlock"]
            # Row 2 should have the spinlock data
            assert sheet.cell(row=2, column=1).value == "TestSpinlock"
            assert sheet.cell(row=2, column=2).value == "STANDARD"
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_os_task_rte_resources(self):
        """
        Test task sheet with RTE-pattern resources.

        Implements: TC_UNIT_REPORTER_00017
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            # Create an application for task mapping.
            # The ref's short name (last path segment) must match the task name.
            app = OsApplication(os_mod, "App1")
            app.setOsTrusted(True)
            app.setOsApplicationCoreAssignment(0)
            app.addOsAppTaskRef(EcucRefType("/Os/TaskRte"))
            os_mod.addOsApplication(app)

            task = OsTask(os_mod, "TaskRte")
            task.setOsTaskPriority(10)
            task.setOsTaskActivation(1)
            task.setOsTaskSchedule("FULL")
            task.setOsStacksize(1024)
            # Add RTE-pattern resources (lines 161-162)
            task.addOsTaskResourceRef(EcucRefType("Rte_CanIf"))
            task.addOsTaskResourceRef(EcucRefType("Rte_Com"))
            # Add non-RTE resource that should be filtered
            task.addOsTaskResourceRef(EcucRefType("Res_Standard"))
            os_mod.addOsTask(task)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsTask" in wb.sheetnames
            sheet = wb["OsTask"]
            # Row 2 should have the task data
            assert sheet.cell(row=2, column=1).value == "TaskRte"
            # Column 9 should contain RTE resources (2 items, wrapped)
            cell_value = sheet.cell(row=2, column=9).value
            assert cell_value is not None
            assert "Rte_CanIf" in cell_value
            assert "Rte_Com" in cell_value
            # Non-RTE resource should not be in the cell
            assert "Res_Standard" not in cell_value
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_os_task_many_rte_resources(self):
        """
        Test task sheet with >10 RTE resources (Total: N summary).

        Implements: TC_UNIT_REPORTER_00017
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            # Create an application for task mapping
            app = OsApplication(os_mod, "App1")
            app.setOsTrusted(True)
            app.setOsApplicationCoreAssignment(0)
            app.addOsAppTaskRef(EcucRefType("/Os/TaskManyRte"))
            os_mod.addOsApplication(app)

            task = OsTask(os_mod, "TaskManyRte")
            task.setOsTaskPriority(5)
            task.setOsTaskActivation(1)
            task.setOsTaskSchedule("FULL")
            task.setOsStacksize(4096)
            # Add 12 RTE-pattern resources (>10 triggers total summary on line 165)
            for i in range(12):
                task.addOsTaskResourceRef(EcucRefType("Rte_Resource_%d" % i))
            os_mod.addOsTask(task)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsTask" in wb.sheetnames
            sheet = wb["OsTask"]
            # Column 9 should show total count summary
            cell_value = sheet.cell(row=2, column=9).value
            assert cell_value is not None
            assert "Total: 12 OsResources" in cell_value
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_os_os_and_hooks(self):
        """
        Test OsOS and OsHooks sheets are produced when data exists.

        Implements: TC_UNIT_REPORTER_00017
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            os_os = OsOS(os_mod, "OsOS")
            os_os.setOsScalabilityClass("SC1")
            os_os.setOsNumberOfCores(2)
            os_os.setOsStackMonitoring(True)
            os_os.setOsUseGetServiceId(False)
            os_os.setOsUseParameterAccess(True)
            os_os.setOsUseResScheduler(True)
            os_os.setOsStatus("STANDARD")
            os_mod.setOsOS(os_os)

            hooks = OsHooks(os_mod, "OsHooks")
            hooks.setOsErrorHook(True)
            hooks.setOsShutdownHook(False)
            hooks.setOsStartupHook(True)
            hooks.setOsPreTaskHook(False)
            hooks.setOsPostTaskHook(True)
            hooks.setOsProtectionHook(False)
            os_mod.setOsHooks(hooks)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsOS" in wb.sheetnames
            assert "OsHooks" in wb.sheetnames

            os_sheet = wb["OsOS"]
            assert os_sheet.cell(row=2, column=1).value == "ScalabilityClass"
            assert os_sheet.cell(row=2, column=2).value == "SC1"
            assert os_sheet.cell(row=3, column=1).value == "NumberOfCores"
            assert os_sheet.cell(row=3, column=2).value == 2

            hooks_sheet = wb["OsHooks"]
            assert hooks_sheet.cell(row=2, column=1).value == "ErrorHook"
            assert hooks_sheet.cell(row=2, column=2).value == "Y"
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_mk_memory_regions(self):
        """
        Test MkMemoryRegion sheet production with microkernel data.

        Implements: TC_UNIT_REPORTER_00017
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            kernel = OsMicrokernel(os_mod, "OsMicrokernel")
            protection = MkMemoryProtection(kernel, "MkMemoryProtection")
            region = MkMemoryRegion(protection, "Region1")
            region.setMkMemoryRegionFlags(1)
            region.setMkMemoryRegionInitialize(True)
            region.setMkMemoryRegionGlobal(False)
            protection.addMkMemoryRegion(region)
            kernel.setMkMemoryProtection(protection)
            os_mod.setOsMicrokernel(kernel)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "MkMemoryRegion" in wb.sheetnames
            sheet = wb["MkMemoryRegion"]
            assert sheet.cell(row=2, column=1).value == "Region1"
            assert sheet.cell(row=2, column=2).value == 1
            assert sheet.cell(row=2, column=3).value == "Y"
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_mk_memory_region_exists_false_paths(self):
        """
        Test mk_memory_region_exists returns False for missing data paths.

        Implements: TC_UNIT_REPORTER_00017
        """
        doc = EBModel.getInstance()
        os_mod = doc.getOs()
        writer = OsXdmXlsWriter()

        # No microkernel set
        assert writer.mk_memory_region_exists(doc) is False

        # Microkernel without protection
        kernel = OsMicrokernel(os_mod, "OsMicrokernel")
        os_mod.setOsMicrokernel(kernel)
        assert writer.mk_memory_region_exists(doc) is False

        # Protection without regions
        protection = MkMemoryProtection(kernel, "MkMemoryProtection")
        kernel.setMkMemoryProtection(protection)
        assert writer.mk_memory_region_exists(doc) is False

        # Protection with regions
        region = MkMemoryRegion(protection, "Region1")
        protection.addMkMemoryRegion(region)
        assert writer.mk_memory_region_exists(doc) is True

    def test_write_appmodes_sheet(self):
        """
        Test Application Modes sheet writing.

        Implements: UTS_OS_REPORTER_00025
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            from eb_model.models.core.os_xdm import OsAppMode
            appmode1 = OsAppMode(os_mod, "AppMode1")
            appmode2 = OsAppMode(os_mod, "AppMode2")
            os_mod.addOsAppMode(appmode1)
            os_mod.addOsAppMode(appmode2)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsAppMode" in wb.sheetnames
            sheet = wb["OsAppMode"]
            assert sheet.cell(row=1, column=1).value == "Name"
            assert sheet.cell(row=2, column=1).value == "AppMode1"
            assert sheet.cell(row=3, column=1).value == "AppMode2"
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_peripheral_areas_sheet(self):
        """
        Test Peripheral Areas sheet writing.

        Implements: UTS_OS_REPORTER_00026
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            from eb_model.models.core.os_xdm import OsPeripheralArea
            area = OsPeripheralArea(os_mod, "PeripheralArea1")
            area.setOsPeripheralAreaStartAddress(4096)
            area.setOsPeripheralAreaEndAddress(8191)
            area.setOsPeripheralAreaId(1)
            area.setOsPeripheralAreaAccessPermission("READ-WRITE")
            os_mod.addOsPeripheralArea(area)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsPeripheralArea" in wb.sheetnames
            sheet = wb["OsPeripheralArea"]
            assert sheet.cell(row=1, column=1).value == "Name"
            assert sheet.cell(row=1, column=2).value == "Start Address"
            assert sheet.cell(row=1, column=3).value == "End Address"
            assert sheet.cell(row=1, column=4).value == "ID"
            assert sheet.cell(row=1, column=5).value == "Access Permission"
            assert sheet.cell(row=2, column=1).value == "PeripheralArea1"
            assert sheet.cell(row=2, column=2).value == 4096
            assert sheet.cell(row=2, column=3).value == 8191
            assert sheet.cell(row=2, column=4).value == 1
            assert sheet.cell(row=2, column=5).value == "READ-WRITE"
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_multiple_os_applications(self):
        """
        Test that multiple OsApplications each appear in a separate row.

        Regression test for missing row += 1 in write_os_applications.

        Implements: UTS_OS_REPORTER_00027
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            app_alpha = OsApplication(os_mod, "AppAlpha")
            app_alpha.setOsTrusted(True)
            app_alpha.setOsApplicationCoreAssignment(0)
            app_alpha.setOsAppEcucPartitionRef(EcucRefType("/Os/PartitionAlpha"))
            os_mod.addOsApplication(app_alpha)

            app_beta = OsApplication(os_mod, "AppBeta")
            app_beta.setOsTrusted(False)
            app_beta.setOsApplicationCoreAssignment(1)
            app_beta.setOsAppEcucPartitionRef(EcucRefType("/Os/PartitionBeta"))
            os_mod.addOsApplication(app_beta)

            app_gamma = OsApplication(os_mod, "AppGamma")
            app_gamma.setOsTrusted(True)
            app_gamma.setOsApplicationCoreAssignment(2)
            os_mod.addOsApplication(app_gamma)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsApplications" in wb.sheetnames
            sheet = wb["OsApplications"]

            assert sheet.cell(row=2, column=1).value == "AppAlpha"
            assert sheet.cell(row=2, column=2).value is True
            assert sheet.cell(row=2, column=3).value == 0
            assert sheet.cell(row=2, column=4).value == "PartitionAlpha"

            assert sheet.cell(row=3, column=1).value == "AppBeta"
            assert sheet.cell(row=3, column=2).value is False
            assert sheet.cell(row=3, column=3).value == 1
            assert sheet.cell(row=3, column=4).value == "PartitionBeta"

            assert sheet.cell(row=4, column=1).value == "AppGamma"
            assert sheet.cell(row=4, column=2).value is True
            assert sheet.cell(row=4, column=3).value == 2

            assert sheet.cell(row=5, column=1).value is None
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_multiple_os_tasks(self):
        """
        Test that multiple OsTasks each appear in a separate row.

        Implements: UTS_OS_REPORTER_00028
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            app = OsApplication(os_mod, "AppForTasks")
            app.addOsAppTaskRef(EcucRefType("/Os/TaskAlpha"))
            app.addOsAppTaskRef(EcucRefType("/Os/TaskBeta"))
            os_mod.addOsApplication(app)

            task_alpha = OsTask(os_mod, "TaskAlpha")
            task_alpha.setOsTaskPriority(1)
            task_alpha.setOsTaskActivation(1)
            task_alpha.setOsTaskSchedule("FULL")
            task_alpha.setOsStacksize(512)
            os_mod.addOsTask(task_alpha)

            task_beta = OsTask(os_mod, "TaskBeta")
            task_beta.setOsTaskPriority(2)
            task_beta.setOsTaskActivation(1)
            task_beta.setOsTaskSchedule("NON")
            task_beta.setOsStacksize(1024)
            os_mod.addOsTask(task_beta)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsTask" in wb.sheetnames
            sheet = wb["OsTask"]

            assert sheet.cell(row=2, column=1).value == "TaskAlpha"
            assert sheet.cell(row=3, column=1).value == "TaskBeta"
            assert sheet.cell(row=4, column=1).value is None
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_multiple_os_isrs(self):
        """
        Test that multiple OsIsrs each appear in a separate row.

        Implements: UTS_OS_REPORTER_00029
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            app = OsApplication(os_mod, "AppForIsrs")
            app.addOsAppIsrRef(EcucRefType("/Os/IsrAlpha"))
            app.addOsAppIsrRef(EcucRefType("/Os/IsrBeta"))
            os_mod.addOsApplication(app)

            isr_alpha = OsIsr(os_mod, "IsrAlpha")
            isr_alpha.setOsIsrCategory("CATEGORY_2")
            isr_alpha.setOsIsrPriority(10)
            os_mod.addOsIsr(isr_alpha)

            isr_beta = OsIsr(os_mod, "IsrBeta")
            isr_beta.setOsIsrCategory("CATEGORY_1")
            isr_beta.setOsIsrPriority(20)
            os_mod.addOsIsr(isr_beta)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsIsr" in wb.sheetnames
            sheet = wb["OsIsr"]

            assert sheet.cell(row=2, column=1).value == "IsrAlpha"
            assert sheet.cell(row=3, column=1).value == "IsrBeta"
            assert sheet.cell(row=4, column=1).value is None
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_multiple_os_spinlocks(self):
        """
        Test that multiple OsSpinlocks each appear in a separate row.

        Implements: UTS_OS_REPORTER_00030
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            spinlock_alpha = OsSpinlock(os_mod, "SpinlockAlpha")
            spinlock_alpha.setOsSpinlockLockMethod("STANDARD")
            os_mod.addOsSpinlock(spinlock_alpha)

            spinlock_beta = OsSpinlock(os_mod, "SpinlockBeta")
            spinlock_beta.setOsSpinlockLockMethod("ALL_INTERRUPT_BLOCKING")
            os_mod.addOsSpinlock(spinlock_beta)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsSpinlock" in wb.sheetnames
            sheet = wb["OsSpinlock"]

            assert sheet.cell(row=2, column=1).value == "SpinlockAlpha"
            assert sheet.cell(row=2, column=2).value == "STANDARD"
            assert sheet.cell(row=3, column=1).value == "SpinlockBeta"
            assert sheet.cell(row=3, column=2).value == "ALL_INTERRUPT_BLOCKING"
            assert sheet.cell(row=4, column=1).value is None
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_multiple_os_counters(self):
        """
        Test that multiple OsCounters each appear in a separate row.

        Implements: UTS_OS_REPORTER_00031
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            counter_alpha = OsCounter(os_mod, "CounterAlpha")
            counter_alpha.setOsCounterMaxAllowedValue(65535)
            counter_alpha.setOsCounterMinCycle(1)
            counter_alpha.setOsCounterTicksPerBase(1)
            counter_alpha.setOsCounterType("SOFTWARE")
            os_mod.addOsCounter(counter_alpha)

            counter_beta = OsCounter(os_mod, "CounterBeta")
            counter_beta.setOsCounterMaxAllowedValue(1000)
            counter_beta.setOsCounterMinCycle(2)
            counter_beta.setOsCounterTicksPerBase(10)
            counter_beta.setOsCounterType("HARDWARE")
            os_mod.addOsCounter(counter_beta)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsCounter" in wb.sheetnames
            sheet = wb["OsCounter"]

            assert sheet.cell(row=2, column=1).value == "CounterAlpha"
            assert sheet.cell(row=2, column=2).value == 65535
            assert sheet.cell(row=3, column=1).value == "CounterBeta"
            assert sheet.cell(row=3, column=2).value == 1000
            assert sheet.cell(row=4, column=1).value is None
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_multiple_os_schedule_tables(self):
        """
        Test that multiple OsScheduleTables each appear in a separate row.

        Implements: UTS_OS_REPORTER_00032
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            counter = OsCounter(os_mod, "OsCounterForTables")
            counter.setOsCounterMaxAllowedValue(65535)
            counter.setOsCounterMinCycle(1)
            counter.setOsCounterTicksPerBase(1)
            os_mod.addOsCounter(counter)

            table_alpha = OsScheduleTable(os_mod, "TableAlpha")
            table_alpha.setOsScheduleTableDuration(100)
            table_alpha.setOsScheduleTableRepeating(True)
            table_alpha.setOsScheduleTableCounterRef(EcucRefType("/Os/OsCounterForTables"))
            os_mod.addOsScheduleTable(table_alpha)

            table_beta = OsScheduleTable(os_mod, "TableBeta")
            table_beta.setOsScheduleTableDuration(200)
            table_beta.setOsScheduleTableRepeating(False)
            table_beta.setOsScheduleTableCounterRef(EcucRefType("/Os/OsCounterForTables"))
            os_mod.addOsScheduleTable(table_beta)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsScheduleTable" in wb.sheetnames
            sheet = wb["OsScheduleTable"]

            assert sheet.cell(row=2, column=1).value == "TableAlpha"
            assert sheet.cell(row=3, column=1).value == "TableBeta"
            assert sheet.cell(row=4, column=1).value is None
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_multiple_mk_memory_regions(self):
        """
        Test that multiple MkMemoryRegions each appear in a separate row.

        Implements: UTS_OS_REPORTER_00033
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            kernel = OsMicrokernel(os_mod, "OsMicrokernelMulti")
            protection = MkMemoryProtection(kernel, "MkMemoryProtectionMulti")

            region_alpha = MkMemoryRegion(protection, "RegionAlpha")
            region_alpha.setMkMemoryRegionFlags(1)
            region_alpha.setMkMemoryRegionInitialize(True)
            region_alpha.setMkMemoryRegionGlobal(False)
            protection.addMkMemoryRegion(region_alpha)

            region_beta = MkMemoryRegion(protection, "RegionBeta")
            region_beta.setMkMemoryRegionFlags(2)
            region_beta.setMkMemoryRegionInitialize(False)
            region_beta.setMkMemoryRegionGlobal(True)
            protection.addMkMemoryRegion(region_beta)

            kernel.setMkMemoryProtection(protection)
            os_mod.setOsMicrokernel(kernel)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "MkMemoryRegion" in wb.sheetnames
            sheet = wb["MkMemoryRegion"]

            assert sheet.cell(row=2, column=1).value == "RegionAlpha"
            assert sheet.cell(row=3, column=1).value == "RegionBeta"
            assert sheet.cell(row=4, column=1).value is None
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)

    def test_write_multiple_os_peripheral_areas(self):
        """
        Test that multiple OsPeripheralAreas each appear in a separate row.

        Implements: UTS_OS_REPORTER_00034
        """
        writer = OsXdmXlsWriter()

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            filename = f.name

        try:
            doc = EBModel.getInstance()
            os_mod = doc.getOs()

            area_alpha = OsPeripheralArea(os_mod, "AreaAlpha")
            area_alpha.setOsPeripheralAreaStartAddress(0x1000)
            area_alpha.setOsPeripheralAreaEndAddress(0x1FFF)
            area_alpha.setOsPeripheralAreaId(1)
            area_alpha.setOsPeripheralAreaAccessPermission("READ-ONLY")
            os_mod.addOsPeripheralArea(area_alpha)

            area_beta = OsPeripheralArea(os_mod, "AreaBeta")
            area_beta.setOsPeripheralAreaStartAddress(0x2000)
            area_beta.setOsPeripheralAreaEndAddress(0x2FFF)
            area_beta.setOsPeripheralAreaId(2)
            area_beta.setOsPeripheralAreaAccessPermission("READ-WRITE")
            os_mod.addOsPeripheralArea(area_beta)

            writer.write(filename, doc)

            wb = load_workbook(filename)
            assert "OsPeripheralArea" in wb.sheetnames
            sheet = wb["OsPeripheralArea"]

            assert sheet.cell(row=2, column=1).value == "AreaAlpha"
            assert sheet.cell(row=2, column=5).value == "READ-ONLY"
            assert sheet.cell(row=3, column=1).value == "AreaBeta"
            assert sheet.cell(row=3, column=5).value == "READ-WRITE"
            assert sheet.cell(row=4, column=1).value is None
            wb.close()
        finally:
            if os.path.exists(filename):
                os.remove(filename)
